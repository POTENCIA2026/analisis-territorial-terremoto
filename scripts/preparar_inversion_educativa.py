#!/usr/bin/env python3
"""Inventario descriptivo municipal de educación. No calcula puntajes ni severidad."""
from collections import defaultdict
import hashlib
import json
from pathlib import Path

from scripts.preparar_matricula_critica import code, count, norm

ROOT = Path(__file__).resolve().parents[1]
SOURCE = 'data/men_sedes_escolares_afectadas_20260921.xlsx'
DAMAGE = {
    'colapso total': 'Colapso total',
    'colapso parcial': 'Colapso parcial',
    'riesgo inminente de colapso': 'Riesgo inminente de colapso',
    'afectacion parcial': 'Afectación parcial',
    'afectacion menor': 'Afectación menor',
    'sin afectacion': 'Sin afectación',
}


def total(rows):
    return None if any(r['enrollment'] is None for r in rows) else sum(r['enrollment'] for r in rows)


def distribution(rows, field):
    groups = defaultdict(list)
    for r in rows:
        groups[r[field]].append(r)
    return [{'label': label, 'sites': len(group), 'enrollment': total(group)}
            for label, group in sorted(groups.items())]


def aggregate(records):
    groups, seen = defaultdict(list), set()
    for r in records:
        if r['site'] in seen:
            raise ValueError(f'Sede duplicada: {r["site"]}')
        seen.add(r['site'])
        groups[r['code']].append(r)
    result = []
    for municipal, rows in sorted(groups.items()):
        affected = [r for r in rows if r['damage'] in DAMAGE.values() and r['damage'] != 'Sin afectación']
        unknown = sum(r['damage'] not in DAMAGE.values() for r in rows)
        enrolled = total(affected)
        rural = total([r for r in affected if r['zone'] == 'rural'])
        unknown_zone = any(r['zone'] not in ('rural', 'urbana') for r in affected)
        flags = sorted({r['territory'] for r in rows})
        result.append({
            'code': municipal, 'reported_sites': len(rows), 'reported_enrollment': total(rows),
            'affected_sites': None if unknown else len(affected),
            'affected_enrollment': None if unknown else enrolled,
            'rural_enrollment': None if unknown or unknown_zone else rural,
            'rural_percent': 100 * rural / enrolled if not unknown and not unknown_zone and enrolled and rural is not None else None,
            'damage': distribution(rows, 'damage'),
            'strategy': distribution(rows, 'strategy'), 'modality': distribution(rows, 'modality'),
            'service': distribution(rows, 'service'),
            'territory': flags[0] if len(flags) == 1 and flags[0] in ('no', 'zomac', 'pdet-zomac', 'pdet') else None,
            'unclassified_sites': unknown,
            'missing_enrollment_sites': sum(r['enrollment'] is None for r in rows),
        })
    return result


def prepare(root=ROOT):
    from openpyxl import load_workbook
    path = root / SOURCE
    reference = {r['code']: r for r in json.loads((root / 'data/poblacion_relativa.json').read_text(encoding='utf-8'))['rows']}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook['Sheet1']
    values = sheet.iter_rows(values_only=True)
    header = next(values)
    if len(header) != 52 or norm(header[9]) != 'codigo dane sede' or norm(header[38]) != 'matricula total nacional':
        raise ValueError('Cambió el esquema MEN: revisar columnas antes de extraer.')
    records = []
    labels = {}
    for row in values:
        if not any(v is not None for v in row):
            continue
        municipal = code(row[3], 5)
        if municipal not in reference or norm(row[1]) != norm(reference[municipal]['d']):
            raise ValueError(f'Municipio sin correspondencia DANE: {municipal}')
        if norm(row[7]) != 'oficial':
            raise ValueError('El alcance oficial del archivo cambió.')
        record = {'code': municipal, 'site': code(row[9], 12), 'enrollment': count(row[38]),
                  'damage': DAMAGE.get(norm(row[41]), 'Sin clasificación'),
                  'zone': norm(row[11]), 'territory': norm(row[43])}
        for field, index in (('service', 49), ('strategy', 50), ('modality', 51)):
            key = norm(row[index])
            # Unifica mayúsculas y espacios, sin interpretar modalidades como presencial/no presencial.
            record[field] = labels.setdefault((field, key), ' '.join(str(row[index]).strip().split()) if key else 'Sin dato')
        records.append(record)
    workbook.close()
    return {'version': 1, 'source': {'file': SOURCE, 'sheet': 'Sheet1', 'file_date': '2026-09-21',
            'date_note': 'Fecha del nombre del archivo; el periodo de la matrícula no está acreditado.',
            'scope': 'Sedes oficiales incluidas en el reporte MEN. No es el universo completo de sedes del municipio.',
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest()},
            'municipalities': aggregate(records)}


if __name__ == '__main__':
    result = prepare()
    target = ROOT / 'data/inversion_educativa.json'
    target.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n', encoding='utf-8')
    print(f'Inventario educativo: {len(result["municipalities"])} municipios. Sin puntajes.')
