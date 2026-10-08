#!/usr/bin/env python3
"""Universo nacional MEN por sede: descarga verificable y resumen sin tasas entre años."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import gzip
import hashlib
import json
from pathlib import Path
import urllib.parse
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
DATASET = 'ngw5-c5nw'
COVERAGE_DATASET = 'nudc-7mev'
FIELDS = 'cod_dane_departamento,departamento,cod_dane_municipio,municipio,codigo_dane_sede,nombre_sede,codigo_dane,nombre_establecimiento,sector,zona'
CONTROL = 'count(*) as filas,count(total_matricula) as filas_con_matricula,sum(total_matricula) as matricula,count(distinct codigo_dane_sede) as sedes,count(distinct cod_dane_municipio) as municipios,max(anno_inf) as ultimo_ano'


def integer(value):
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n < 0 or n != n.to_integral_value():
            raise ValueError('Se esperaba un entero no negativo.')
        return int(n)
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f'Número no válido: {value!r}') from exc


def dane(value, length):
    raw = str(integer(value))
    if len(raw) > length or (length == 12 and len(raw) != 12):
        raise ValueError(f'Código DANE de longitud incorrecta: {value!r}')
    return raw.zfill(length)


def fetch(dataset, params=None, metadata=False):
    url = f'https://www.datos.gov.co/api/views/{dataset}.json' if metadata else (
        f'https://www.datos.gov.co/resource/{dataset}.json?' + urllib.parse.urlencode(params or {}))
    request = urllib.request.Request(url, headers={'User-Agent': 'Infopotencia-MEN/1.0'})
    with urllib.request.urlopen(request, timeout=240) as response:
        return json.load(response)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(data, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n'
    path.write_bytes(gzip.compress(text.encode(), mtime=0)) if path.suffix == '.gz' else path.write_text(text, encoding='utf-8')


def read_json(path):
    return json.loads(gzip.decompress(path.read_bytes()) if path.suffix == '.gz' else path.read_bytes())


def download(target, year=2025, coverage_year=2024):
    before = fetch(DATASET, metadata=True)
    coverage_before = fetch(COVERAGE_DATASET, metadata=True)
    rows, offset = [], 0
    params = {'$select': FIELDS + ',sum(total_matricula) as matricula,count(*) as filas,count(total_matricula) as filas_con_matricula',
              '$where': f'anno_inf={year}', '$group': FIELDS, '$order': FIELDS, '$limit': '50000'}
    while True:
        batch = fetch(DATASET, dict(params, **{'$offset': str(offset)}))
        rows.extend(batch)
        print(f'MEN {year}: {len(rows)} grupos por sede descargados.', flush=True)
        if len(batch) < 50000:
            break
        offset += len(batch)
    control = fetch(DATASET, {'$select': CONTROL, '$where': f'anno_inf={year}'})[0]
    coverage = fetch(COVERAGE_DATASET, {'$where': f"a_o='{coverage_year}'", '$order': 'c_digo_municipio', '$limit': '50000'})
    coverage_count = fetch(COVERAGE_DATASET, {'$select': 'count(*) as n', '$where': f"a_o='{coverage_year}'"})[0]
    if len(coverage) != integer(coverage_count['n']):
        raise ValueError('La descarga de cobertura está incompleta.')
    after = fetch(DATASET, metadata=True)
    coverage_after = fetch(COVERAGE_DATASET, metadata=True)
    if before['rowsUpdatedAt'] != after['rowsUpdatedAt'] or coverage_before['rowsUpdatedAt'] != coverage_after['rowsUpdatedAt']:
        raise ValueError('La fuente cambió durante la descarga. Repetir para obtener un único corte.')
    # No guardar una extracción que no concilie con controles independientes del servidor.
    summarize_sites(rows, control)
    manifest = {'downloaded_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                'enrollment_year': year, 'coverage_year': coverage_year, 'enrollment_dataset': DATASET,
                'coverage_dataset': COVERAGE_DATASET, 'enrollment_query': params,
                'enrollment_rows_updated_at': before['rowsUpdatedAt'],
                'coverage_rows_updated_at': coverage_before['rowsUpdatedAt'], 'control': control}
    for name, data in [('sedes.json.gz', rows), ('cobertura.json', coverage), ('fuentes.json', manifest)]:
        write_json(target / name, data)


def summarize_sites(rows, control):
    sites, groups = {}, defaultdict(list)
    source_rows = 0
    for row in rows:
        site = dane(row['codigo_dane_sede'], 12)
        code = dane(row['cod_dane_municipio'], 5)
        if site in sites:
            raise ValueError(f'Sede duplicada o con atributos incompatibles: {site}')
        if not code.startswith(dane(row['cod_dane_departamento'], 2)):
            raise ValueError(f'Departamento inconsistente: {code}')
        if row['sector'] not in ('OFICIAL', 'NO_OFICIAL') or row['zona'] not in ('Urbana', 'Rural'):
            raise ValueError(f'Sector o zona desconocido: {site}')
        if integer(row['filas']) != integer(row['filas_con_matricula']):
            raise ValueError(f'Matrícula incompleta: {site}')
        enrollment = integer(row['matricula'])
        source_rows += integer(row['filas'])
        entry = {'code': code, 'site': site, 'sector': row['sector'], 'zone': row['zona'], 'enrollment': enrollment}
        sites[site] = entry
        groups[code].append(entry)
    totals = {'filas': source_rows, 'filas_con_matricula': source_rows, 'sedes': len(sites),
              'municipios': len(groups), 'matricula': sum(s['enrollment'] for s in sites.values())}
    for key, actual in totals.items():
        if actual != integer(control[key]):
            raise ValueError(f'No concilia {key}: {actual} != {control[key]}')
    result = []
    for code, entries in sorted(groups.items()):
        sectors = [{'label': label, 'sites': sum(s['sector'] == key for s in entries),
                    'enrollment': sum(s['enrollment'] for s in entries if s['sector'] == key)}
                   for key, label in [('OFICIAL', 'Oficial'), ('NO_OFICIAL', 'No oficial')]]
        zones = [{'label': zone, 'sites': sum(s['zone'] == zone for s in entries),
                  'enrollment': sum(s['enrollment'] for s in entries if s['zone'] == zone)} for zone in ('Rural', 'Urbana')]
        enrollment = sum(s['enrollment'] for s in entries)
        result.append({'code': code, 'sites': len(entries), 'enrollment': enrollment, 'sectors': sectors, 'zones': zones,
                       'official_percent': 100 * sectors[0]['enrollment'] / enrollment if enrollment else None,
                       'rural_percent': 100 * zones[0]['enrollment'] / enrollment if enrollment else None})
    return result, sites, totals


def summarize_coverage(rows, year):
    result, seen = [], set()
    for row in rows:
        code = dane(row['c_digo_municipio'], 5)
        if code in seen or integer(row['a_o']) != year:
            raise ValueError(f'Cobertura duplicada o de otro año: {code}')
        seen.add(code)
        raw = row.get('cobertura_neta')
        value = None if raw in (None, '') else float(Decimal(str(raw)))
        if value is not None and (not Decimal(str(value)).is_finite() or value < 0):
            raise ValueError(f'Cobertura no válida: {code}')
        # El API publica puntos porcentuales: 88.76 significa 88,76 %, no 0,8876 %.
        result.append({'code': code, 'year': year, 'net_coverage': value})
    return sorted(result, key=lambda r: r['code'])


def crosswalk(sites, reported):
    groups, seen, exceptions = defaultdict(Counter), set(), []
    for row in reported:
        site, code = dane(row['site'], 12), dane(row['code'], 5)
        if site in seen:
            raise ValueError(f'Sede duplicada en reporte de afectadas: {site}')
        seen.add(site)
        match = sites.get(site)
        status = 'missing' if match is None else ('municipality_conflict' if match['code'] != code else ('sector_conflict' if match['sector'] != 'OFICIAL' else 'matched'))
        groups[code]['reported'] += 1
        groups[code][status] += 1
        if status != 'matched':
            exceptions.append({'site': site, 'code': code, 'status': status, 'national_code': match['code'] if match else None})
    return [{'code': code, **{key: counts[key] for key in ('reported', 'matched', 'missing', 'municipality_conflict', 'sector_conflict')}}
            for code, counts in sorted(groups.items())], exceptions


def reported_sites(root):
    from openpyxl import load_workbook
    path = root / 'data/men_sedes_escolares_afectadas_20260921.xlsx'
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = workbook['Sheet1'].iter_rows(values_only=True)
    header = next(values)
    if len(header) != 52 or str(header[9]).strip().upper() != 'CODIGO DANE SEDE':
        workbook.close()
        raise ValueError('Cambió el esquema del reporte de sedes afectadas.')
    rows = [{'code': row[3], 'site': row[9]} for row in values if any(v is not None for v in row)]
    workbook.close()
    return rows, hashlib.sha256(path.read_bytes()).hexdigest()


def prepare(root=ROOT):
    source = root / 'data/men_abierto'
    manifest = read_json(source / 'fuentes.json')
    if integer(manifest['control']['ultimo_ano']) != manifest['enrollment_year']:
        raise ValueError('El año del control no corresponde al año de matrícula seleccionado.')
    municipalities, sites, totals = summarize_sites(read_json(source / 'sedes.json.gz'), manifest['control'])
    coverage = summarize_coverage(read_json(source / 'cobertura.json'), manifest['coverage_year'])
    reported, report_hash = reported_sites(root)
    matches, exceptions = crosswalk(sites, reported)
    roster = {r['code'] for r in read_json(root / 'data/poblacion_relativa.json')['rows']}
    unknown = sorted({r['code'] for r in municipalities + coverage} - roster)
    if unknown:
        raise ValueError(f'Municipios fuera del catálogo del tablero: {unknown}')
    manifest = dict(manifest, hashes={name: hashlib.sha256((source / name).read_bytes()).hexdigest()
                                    for name in ('sedes.json.gz', 'cobertura.json')}, affected_report_sha256=report_hash)
    return {'version': 1, 'source': manifest, 'totals': totals, 'municipalities': municipalities,
            'coverage': coverage, 'crosswalk': matches, 'crosswalk_exceptions': exceptions,
            'scope': 'Sedes con matrícula reportada en el conjunto nacional del MEN para el año seleccionado. No es un censo de sedes sin matrícula ni de cupos disponibles.'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download', action='store_true', help='Actualizar instantáneas desde datos.gov.co; sin esta opción se trabaja sin red.')
    parser.add_argument('--year', type=int, default=2025)
    parser.add_argument('--coverage-year', type=int, default=2024)
    args = parser.parse_args()
    if args.download:
        download(ROOT / 'data/men_abierto', args.year, args.coverage_year)
    payload = prepare()
    write_json(ROOT / 'data/men_universo.json', payload)
    print(json.dumps({'totals': payload['totals'], 'crosswalk': dict(sum((Counter({k: r[k] for k in ('reported','matched','missing','municipality_conflict','sector_conflict')}) for r in payload['crosswalk']), Counter()))}, ensure_ascii=False))
